# -*- coding: utf-8 -*-
"""
md_to_word.py —— Markdown 转 Word(.docx)
========================================
纯 Python 实现（python-docx），不依赖 Microsoft Word / WPS / LibreOffice / COM 自动化，
可直接随 PyInstaller 打包分发。

支持的 Markdown 语法：
    标题 #~######、普通段落（自动合并软换行）、无序/有序列表（含两级嵌套）、任务列表、
    Markdown 管道表格、围栏代码块、引用块、分隔线、
    行内加粗/斜体/删除线/行内代码、链接、图片（本地相对路径，自动缩放）

用法：
    from md_to_word import convert_md_file, convert_md_directory

    convert_md_file("a.md")            # 同目录生成 a.docx
    convert_md_directory("汇总MD")      # 目录内每个 .md 各生成一份 .docx
"""

from __future__ import annotations

import os
import re
import tempfile
from urllib.parse import unquote

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.image.image import Image as _DocxImage
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

# ---------------------------------------------------------------- 版式常量

BODY_ASCII_FONT = "Times New Roman"
BODY_EAST_ASIA_FONT = "宋体"
HEADING_EAST_ASIA_FONT = "微软雅黑"
CODE_FONT = "Consolas"

BODY_FONT_SIZE = Pt(10.5)  # 五号
CODE_FONT_SIZE = Pt(9)
HEADING_SIZES = {1: Pt(18), 2: Pt(16), 3: Pt(14), 4: Pt(12), 5: Pt(11), 6: Pt(10.5)}

MAX_IMAGE_WIDTH_IN = 6.0
MAX_IMAGE_HEIGHT_IN = 8.5

# ---------------------------------------------------------------- 正则

_INLINE_RE = re.compile(
    r"(!\[[^\]]*\]\([^)]*\))"  # 1 图片
    r"|(\*\*.+?\*\*)"  # 2 粗体
    r"|(__.+?__)"  # 3 粗体
    r"|(~~.+?~~)"  # 4 删除线
    r"|(`+[^`]+`+)"  # 5 行内代码
    r"|(\*[^*\n]+?\*)"  # 6 斜体
    r"|(_[^_\n]+?_)"  # 7 斜体
    r"|(\[[^\]]+\]\([^)]+\))"  # 8 链接
)

_HTML_TAG_RE = re.compile(r"</?[A-Za-z][^>]*>")
_FORMULA_RE = re.compile(r"(\$\$[^$]+\$\$|\$[^$\n]{1,500}\$)")
_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
_HR_RE = re.compile(r"^\s{0,3}([-*_])(\s*\1){2,}\s*$")
_FENCE_RE = re.compile(r"^\s*(```|~~~)")
_LIST_RE = re.compile(r"^(\s*)([-*+]|\d+[.)])\s+(.*)$")
_TASK_RE = re.compile(r"^\[([ xX])\]\s*(.*)$")
_SEP_CELL_RE = re.compile(r"^:?-{2,}:?$")
_BR_RE = re.compile(r"(?i)<br\s*/?>")

_HTML_ENTITIES = (
    ("&nbsp;", " "),
    ("&emsp;", "\u3000"),
    ("&ensp;", " "),
    ("&quot;", '"'),
    ("&#39;", "'"),
    ("&lt;", "<"),
    ("&gt;", ">"),
    ("&amp;", "&"),
)


# ---------------------------------------------------------------- 低层工具


def _clean_html(text: str) -> str:
    """剥离残留的 HTML 标签并反转义常见实体。"""
    if not text:
        return ""
    text = _HTML_TAG_RE.sub("", text)
    for src, dst in _HTML_ENTITIES:
        text = text.replace(src, dst)
    return text


def _apply_run_font(
    run,
    ascii_font=BODY_ASCII_FONT,
    east_asia_font=BODY_EAST_ASIA_FONT,
    size=None,
    bold=None,
    italic=None,
    underline=None,
    color=None,
):
    """统一设置 run 的西文/中文字体、字号与字形。"""
    run.font.name = ascii_font
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), ascii_font)
    rFonts.set(qn("w:hAnsi"), ascii_font)
    rFonts.set(qn("w:eastAsia"), east_asia_font)
    if size is not None:
        run.font.size = size
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if underline is not None:
        run.font.underline = underline
    if color is not None:
        run.font.color.rgb = color


def _shade_element(element, fill="F2F2F2"):
    """给段落或 run 的 XML 节点加底纹。"""
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    element.append(shd)


def _add_horizontal_rule(doc):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "AAAAAA")
    pBdr.append(bottom)
    pPr.append(pBdr)


def _join_soft_lines(lines):
    """Markdown 软换行合并：中文之间不补空格，其余补一个空格。"""
    result = ""
    for index, line in enumerate(lines):
        if index == 0:
            result = line
            continue
        prev_tail = result[-1:] if result else ""
        next_head = line[:1]
        if _is_cjk(prev_tail) and _is_cjk(next_head):
            result += line
        else:
            result += " " + line
    return result


def _is_cjk(char):
    if not char:
        return False
    code = ord(char)
    return (
        0x3000 <= code <= 0x303F
        or 0x4E00 <= code <= 0x9FFF
        or 0xFF00 <= code <= 0xFFEF
    )


# ---------------------------------------------------------------- 图片


def _resolve_image_path(raw_path, base_dir):
    """把 MD 里的图片链接解析为本地绝对路径，找不到返回 None。"""
    path = (raw_path or "").strip()
    if path.startswith("<") and path.endswith(">"):
        path = path[1:-1].strip()
    # 去掉 Markdown 的可选 title：  images/a.png "标题"
    title_match = re.match(r'^(.*?)\s+["\'].*["\']$', path)
    if title_match:
        path = title_match.group(1).strip()
    path = unquote(path)
    if not path:
        return None
    if re.match(r"^[a-zA-Z][a-zA-Z0-9+.-]*://", path) or path.startswith("data:"):
        return None
    candidate = path if os.path.isabs(path) else os.path.join(base_dir, path)
    candidate = os.path.normpath(candidate)
    if os.path.isfile(candidate):
        return candidate
    # 退化匹配：仅用文件名在 base_dir 附近查找
    file_name = os.path.basename(path.replace("\\", "/"))
    if file_name:
        for root, _dirs, files in os.walk(base_dir):
            if file_name in files:
                return os.path.join(root, file_name)
    return None


def _prepare_image(path, tmp_dir):
    """返回可用于插入的图片路径；python-docx 不支持的格式用 PIL 转 PNG。"""
    try:
        _DocxImage.from_file(path)
        return path
    except Exception:
        pass
    try:
        from PIL import Image as _PILImage

        target = os.path.join(tmp_dir, f"{abs(hash(path))}.png")
        with _PILImage.open(path) as image:
            image.convert("RGB").save(target, "PNG")
        return target
    except Exception:
        return None


def _image_display_size_inches(path):
    """按图片自身 DPI 推算显示尺寸（英寸）。"""
    try:
        info = _DocxImage.from_file(path)
        dpi_x = info.horz_dpi if info.horz_dpi and info.horz_dpi > 1 else 96.0
        dpi_y = info.vert_dpi if info.vert_dpi and info.vert_dpi > 1 else 96.0
        width = info.px_width / float(dpi_x)
        height = info.px_height / float(dpi_y)
        return width, height
    except Exception:
        return None, None


def _insert_image(paragraph, token, base_dir, tmp_dir, log):
    match = re.match(r"!\[([^\]]*)\]\(([^)]*)\)", token)
    if not match:
        return
    raw_path = match.group(2)
    resolved = _resolve_image_path(raw_path, base_dir)
    if not resolved:
        run = paragraph.add_run(f"[图片缺失: {raw_path.strip()}]")
        _apply_run_font(run, size=Pt(9), italic=True, color=RGBColor(0xC0, 0x00, 0x00))
        return

    usable = _prepare_image(resolved, tmp_dir)
    if not usable:
        run = paragraph.add_run(f"[图片无法读取: {os.path.basename(resolved)}]")
        _apply_run_font(run, size=Pt(9), italic=True, color=RGBColor(0xC0, 0x00, 0x00))
        return

    width_in, height_in = _image_display_size_inches(usable)
    try:
        if width_in and height_in and width_in > 0 and height_in > 0:
            scale = 1.0
            if width_in > MAX_IMAGE_WIDTH_IN:
                scale = min(scale, MAX_IMAGE_WIDTH_IN / width_in)
            if height_in > MAX_IMAGE_HEIGHT_IN:
                scale = min(scale, MAX_IMAGE_HEIGHT_IN / height_in)
            paragraph.add_run().add_picture(
                usable, width=Inches(width_in * scale), height=Inches(height_in * scale)
            )
        else:
            paragraph.add_run().add_picture(usable, width=Inches(MAX_IMAGE_WIDTH_IN))
        if paragraph.text.strip() == "":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    except Exception as exc:
        log(f"⚠️ 插入图片失败：{os.path.basename(resolved)}，错误：{exc}")


# ---------------------------------------------------------------- 行内内容


def _add_inline_segment(
    paragraph,
    text,
    base_dir,
    tmp_dir,
    log,
    ascii_font,
    east_asia_font,
    size,
    bold,
    italic,
    color,
):
    """处理不含 LaTeX 公式的一段文本。"""
    for seg_index, segment in enumerate(_BR_RE.split(text)):
        if seg_index > 0:
            line_break = paragraph.add_run()
            _apply_run_font(line_break, ascii_font, east_asia_font, size=size, bold=bold, italic=italic)
            line_break.add_break()

        cursor = 0
        for match in _INLINE_RE.finditer(segment):
            if match.start() > cursor:
                prefix = _clean_html(segment[cursor : match.start()])
                if prefix:
                    run = paragraph.add_run(prefix)
                    _apply_run_font(run, ascii_font, east_asia_font, size=size, bold=bold, italic=italic, color=color)
            token = match.group(0)
            if match.group(1):
                _insert_image(paragraph, token, base_dir, tmp_dir, log)
            elif match.group(2) or match.group(3):
                run = paragraph.add_run(_clean_html(token[2:-2]))
                _apply_run_font(run, ascii_font, east_asia_font, size=size, bold=True, italic=italic, color=color)
            elif match.group(4):
                run = paragraph.add_run(_clean_html(token[2:-2]))
                _apply_run_font(run, ascii_font, east_asia_font, size=size, bold=bold, italic=italic, color=color)
                run.font.strike = True
            elif match.group(5):
                run = paragraph.add_run(token.strip("`"))
                _apply_run_font(run, CODE_FONT, CODE_FONT, size=Pt(9.5), color=RGBColor(0xA3, 0x15, 0x15))
                _shade_element(run._element.get_or_add_rPr())
            elif match.group(6) or match.group(7):
                run = paragraph.add_run(_clean_html(token[1:-1]))
                _apply_run_font(run, ascii_font, east_asia_font, size=size, bold=bold, italic=True, color=color)
            elif match.group(8):
                link = re.match(r"\[([^\]]+)\]\(([^)]*)\)", token)
                label = _clean_html(link.group(1)) if link else token
                url = (link.group(2) or "").strip() if link else ""
                run = paragraph.add_run(label or url)
                _apply_run_font(
                    run, ascii_font, east_asia_font, size=size, bold=bold, italic=italic,
                    underline=True, color=RGBColor(0x1F, 0x4E, 0x79),
                )
                if url and url != label:
                    tail = paragraph.add_run(f"（{url}）")
                    _apply_run_font(tail, ascii_font, east_asia_font, size=Pt(8.5), color=RGBColor(0x80, 0x80, 0x80))
            cursor = match.end()

        if cursor < len(segment):
            tail_text = _clean_html(segment[cursor:])
            if tail_text:
                run = paragraph.add_run(tail_text)
                _apply_run_font(run, ascii_font, east_asia_font, size=size, bold=bold, italic=italic, color=color)


def _add_inline(
    paragraph,
    text,
    base_dir,
    tmp_dir,
    log,
    ascii_font=BODY_ASCII_FONT,
    east_asia_font=BODY_EAST_ASIA_FONT,
    size=None,
    bold=False,
    italic=False,
    color=None,
):
    """把一段 Markdown 文本按行内语法写入段落（LaTeX 公式原样保留）。"""
    if text is None:
        return
    text = text.replace("\u200b", "")
    if not text:
        return

    for part_index, part in enumerate(_FORMULA_RE.split(text)):
        if not part:
            continue
        if part_index % 2 == 1:
            run = paragraph.add_run(part)
            _apply_run_font(run, ascii_font, east_asia_font, size=size, bold=bold, italic=italic, color=color)
            continue
        _add_inline_segment(
            paragraph, part, base_dir, tmp_dir, log, ascii_font, east_asia_font, size, bold, italic, color
        )


# ---------------------------------------------------------------- 块级元素


def _add_heading(doc, text, level, base_dir, tmp_dir, log):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(14 if level <= 2 else 10)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_with_next = True
    _add_inline(
        paragraph,
        text,
        base_dir,
        tmp_dir,
        log,
        east_asia_font=HEADING_EAST_ASIA_FONT,
        size=HEADING_SIZES.get(level, BODY_FONT_SIZE),
        bold=True,
    )


def _add_paragraph(doc, text, base_dir, tmp_dir, log):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.35
    _add_inline(paragraph, text, base_dir, tmp_dir, log)


def _add_quote(doc, text, base_dir, tmp_dir, log):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.75)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run("▍")
    _apply_run_font(run, BODY_ASCII_FONT, BODY_EAST_ASIA_FONT, color=RGBColor(0xAA, 0xAA, 0xAA))
    _add_inline(
        paragraph,
        text,
        base_dir,
        tmp_dir,
        log,
        size=Pt(10),
        italic=True,
        color=RGBColor(0x50, 0x50, 0x50),
    )


def _add_code_block(doc, code_lines):
    if not code_lines:
        return
    for raw in code_lines:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.left_indent = Cm(0.4)
        run = paragraph.add_run(raw.rstrip() if raw.strip() else " ")
        _apply_run_font(run, CODE_FONT, CODE_FONT, size=CODE_FONT_SIZE)
        _shade_element(paragraph._p.get_or_add_pPr())


def _add_list_item(doc, text, ordered, level, base_dir, tmp_dir, log):
    style_candidates = (
        ["List Number", "List Number 2", "List Number 3"]
        if ordered
        else ["List Bullet", "List Bullet 2", "List Bullet 3"]
    )
    style_name = style_candidates[min(level, len(style_candidates) - 1)]
    try:
        paragraph = doc.add_paragraph(style=style_name)
    except KeyError:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.75 + 0.6 * level)

    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.3

    task = _TASK_RE.match(text)
    if task:
        box = "☑" if task.group(1).lower() == "x" else "☐"
        run = paragraph.add_run(f"{box} ")
        _apply_run_font(run)
        text = task.group(2)

    _add_inline(paragraph, text, base_dir, tmp_dir, log)


def _add_table(doc, table_lines, base_dir, tmp_dir, log):
    rows = []
    for line in table_lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        rows.append(cells)
    if not rows:
        return

    header = None
    body = rows
    if len(rows) >= 2 and rows[1] and all(_SEP_CELL_RE.match(cell.replace(" ", "")) for cell in rows[1] if cell != ""):
        header = rows[0]
        body = rows[2:]
    elif len(rows) >= 2 and all(set(cell.replace(" ", "")) <= {"-", ":"} and cell.replace(" ", "") for cell in rows[1]):
        header = rows[0]
        body = rows[2:]

    column_count = max(len(row) for row in rows)
    table = doc.add_table(rows=0, cols=column_count)
    table.style = "Table Grid"

    def fill_row(cells, bold=False):
        row_cells = table.add_row().cells
        for index in range(column_count):
            value = _clean_html(cells[index]) if index < len(cells) else ""
            cell_paragraph = row_cells[index].paragraphs[0]
            cell_paragraph.paragraph_format.space_after = Pt(0)
            _add_inline(
                cell_paragraph,
                value,
                base_dir,
                tmp_dir,
                log,
                size=Pt(9.5),
                bold=bold,
            )

    if header:
        fill_row(header, bold=True)
    for data_row in body:
        if not any(cell.strip() for cell in data_row):
            continue
        fill_row(data_row)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(6)


def _parse_html_table(html):
    """解析 HTML <table>，返回 [[{text, colspan, header}, ...], ...]。"""
    rows = []
    for row_match in re.finditer(r"<tr[^>]*>(.*?)</tr>", html, re.S | re.I):
        cells = []
        for cell_match in re.finditer(r"<(td|th)([^>]*)>(.*?)</\1\s*>", row_match.group(1), re.S | re.I):
            attrs = cell_match.group(2) or ""
            span_match = re.search(r'colspan\s*=\s*["\']?(\d+)', attrs, re.I)
            span = int(span_match.group(1)) if span_match else 1
            cells.append(
                {
                    "text": cell_match.group(3),
                    "colspan": max(span, 1),
                    "header": cell_match.group(1).lower() == "th",
                }
            )
        if cells:
            rows.append(cells)
    return rows


def _fill_table_cell(cell, text, base_dir, tmp_dir, log, bold=False):
    """清空单元格并写入内容（保留单元格内换行）。"""
    text = re.sub(r"(?i)<br\s*/?>", "\n", text or "")
    text = _clean_html(text)
    lines = [line.strip() for line in text.split("\n")]
    while lines and not lines[0]:
        lines.pop(0)
    while lines and not lines[-1]:
        lines.pop()
    if not lines:
        lines = [""]

    while len(cell.paragraphs) > 1:
        element = cell.paragraphs[-1]._element
        element.getparent().remove(element)
    first_paragraph = cell.paragraphs[0]
    for run in list(first_paragraph.runs):
        run._element.getparent().remove(run._element)

    for index, line in enumerate(lines):
        paragraph = first_paragraph if index == 0 else cell.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.15
        _add_inline(paragraph, line, base_dir, tmp_dir, log, size=Pt(9), bold=bold)


def _add_html_table(doc, html, base_dir, tmp_dir, log):
    """把 HTML 表格写入 docx，支持 colspan 合并。"""
    rows = _parse_html_table(html)
    if not rows:
        # 无法解析时退化为纯文本，避免内容丢失
        _add_paragraph(doc, _clean_html(re.sub(r"<[^>]+>", " ", html)), base_dir, tmp_dir, log)
        return

    column_count = max(sum(cell["colspan"] for cell in row) for row in rows)
    column_count = min(column_count, 16) if column_count else 1
    table = doc.add_table(rows=len(rows), cols=column_count)
    table.style = "Table Grid"

    for row_index, row in enumerate(rows):
        column = 0
        for cell_spec in row:
            if column >= column_count:
                break
            span = min(cell_spec["colspan"], column_count - column)
            cell = table.cell(row_index, column)
            if span > 1:
                cell = cell.merge(table.cell(row_index, column + span - 1))
            _fill_table_cell(cell, cell_spec["text"], base_dir, tmp_dir, log, bold=cell_spec["header"])
            column += span

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(6)


def _is_block_start(line):
    stripped = line.strip()
    if not stripped:
        return True
    return bool(
        _FENCE_RE.match(line)
        or _HEADING_RE.match(stripped)
        or _HR_RE.match(line)
        or _LIST_RE.match(line)
        or stripped.startswith("|")
        or stripped.startswith(">")
        or stripped.startswith("<!--")
        or "<table" in line.lower()
    )


def _convert_lines(lines, doc, base_dir, tmp_dir, log):
    index = 0
    total = len(lines)
    while index < total:
        raw = lines[index]
        stripped = raw.strip()

        if not stripped:
            index += 1
            continue

        # HTML 注释
        if stripped.startswith("<!--"):
            while index < total and "-->" not in lines[index]:
                index += 1
            index += 1
            continue

        # 围栏代码块
        fence = _FENCE_RE.match(raw)
        if fence:
            marker = fence.group(1)
            index += 1
            code_lines = []
            while index < total and not lines[index].strip().startswith(marker):
                code_lines.append(lines[index])
                index += 1
            index += 1
            _add_code_block(doc, code_lines)
            continue

        # HTML 表格（MinerU 的 full.md 常直接内嵌 <table>）
        lower_raw = raw.lower()
        if "<table" in lower_raw:
            start = lower_raw.index("<table")
            prefix = raw[:start].strip()
            buffer = [raw[start:]]
            index += 1
            while "</table>" not in "\n".join(buffer).lower() and index < total:
                buffer.append(lines[index])
                index += 1
            html = "\n".join(buffer)
            close_index = html.lower().find("</table>")
            if close_index == -1:
                table_html, tail = html, ""
            else:
                close_index += len("</table>")
                table_html, tail = html[:close_index], html[close_index:].strip()
            if prefix:
                _add_paragraph(doc, prefix, base_dir, tmp_dir, log)
            _add_html_table(doc, table_html, base_dir, tmp_dir, log)
            if tail:
                _add_paragraph(doc, tail, base_dir, tmp_dir, log)
            continue

        # Markdown 管道表格
        if stripped.startswith("|"):
            table_lines = []
            while index < total and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            _add_table(doc, table_lines, base_dir, tmp_dir, log)
            continue

        # 分隔线
        if _HR_RE.match(raw):
            _add_horizontal_rule(doc)
            index += 1
            continue

        # 标题
        heading = _HEADING_RE.match(stripped)
        if heading:
            _add_heading(doc, heading.group(2).strip(), len(heading.group(1)), base_dir, tmp_dir, log)
            index += 1
            continue

        # 引用块
        if stripped.startswith(">"):
            quote_lines = []
            while index < total and lines[index].strip().startswith(">"):
                quote_lines.append(lines[index].strip().lstrip(">").strip())
                index += 1
            _add_quote(doc, _join_soft_lines(quote_lines), base_dir, tmp_dir, log)
            continue

        # 列表（支持连续多项与两级嵌套）
        list_match = _LIST_RE.match(raw)
        if list_match:
            while index < total:
                current = lines[index]
                item = _LIST_RE.match(current)
                if not item:
                    break
                indent_raw = item.group(1).replace("\t", "    ")
                level = min(len(indent_raw) // 2, 2)
                ordered = item.group(2)[0].isdigit()
                _add_list_item(doc, item.group(3).strip(), ordered, level, base_dir, tmp_dir, log)
                index += 1
            continue

        # 普通段落：合并后续软换行
        buffer_lines = [stripped]
        index += 1
        while index < total and not _is_block_start(lines[index]):
            buffer_lines.append(lines[index].strip())
            index += 1
        _add_paragraph(doc, _join_soft_lines(buffer_lines), base_dir, tmp_dir, log)


# ---------------------------------------------------------------- 文档骨架


def _create_document():
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = BODY_ASCII_FONT
    normal.font.size = BODY_FONT_SIZE
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), BODY_ASCII_FONT)
    rFonts.set(qn("w:hAnsi"), BODY_ASCII_FONT)
    rFonts.set(qn("w:eastAsia"), BODY_EAST_ASIA_FONT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.35

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    return doc


# ---------------------------------------------------------------- 对外接口


def convert_md_file(md_path, docx_path=None, log=print):
    """把单个 Markdown 文件转换为 .docx，返回生成的文件路径（失败返回 None）。"""
    md_path = os.path.abspath(md_path)
    if not os.path.isfile(md_path):
        log(f"❌ MD 文件不存在：{md_path}")
        return None
    if docx_path is None:
        docx_path = os.path.splitext(md_path)[0] + ".docx"
    docx_path = os.path.abspath(docx_path)

    base_dir = os.path.dirname(md_path)
    try:
        with open(md_path, "r", encoding="utf-8", errors="ignore") as handle:
            lines = handle.read().splitlines()
    except Exception as exc:
        log(f"❌ 读取 MD 失败：{md_path}，错误：{exc}")
        return None

    try:
        doc = _create_document()
        with tempfile.TemporaryDirectory(prefix="md2docx_") as tmp_dir:
            _convert_lines(lines, doc, base_dir, tmp_dir, log)
        doc.save(docx_path)
    except Exception as exc:
        log(f"❌ 转换失败：{os.path.basename(md_path)}，错误：{exc}")
        return None

    log(f"✅ 已生成 Word：{os.path.basename(docx_path)}")
    return docx_path


def convert_md_directory(md_dir, log=print):
    """把目录下所有 .md 批量转换为 .docx，返回生成的文件路径列表。"""
    md_dir = os.path.abspath(md_dir)
    if not os.path.isdir(md_dir):
        log(f"❌ 目录不存在：{md_dir}")
        return []

    md_files = []
    for dirpath, dirnames, filenames in os.walk(md_dir):
        dirnames[:] = [name for name in dirnames if not name.lower().endswith("_assets")]
        for file_name in sorted(filenames):
            if file_name.lower().endswith(".md"):
                md_files.append(os.path.join(dirpath, file_name))

    if not md_files:
        log("⚠️ 未找到可转换的 MD 文件")
        return []

    log(f"共发现 {len(md_files)} 个 MD 文件，开始转换为 Word 文档 ...")
    results = []
    for md_path in md_files:
        docx_path = convert_md_file(md_path, log=log)
        if docx_path:
            results.append(docx_path)

    log(f"🎉 Word 转换完成：成功 {len(results)} / {len(md_files)} 个")
    return results


if __name__ == "__main__":
    import sys

    if len(sys.argv) > 1:
        target = sys.argv[1]
        if os.path.isdir(target):
            convert_md_directory(target)
        else:
            convert_md_file(target)
    else:
        print("用法：python md_to_word.py <md文件或目录>")
